from pathlib import Path

import docx
import pytest

from humanizador.cli import main
from humanizador.lector import Bloque, frases, leer
from humanizador.marcar import marcar_docx, reporte_html
from humanizador.reglas import cargar_perfil, perfiles_disponibles, revisar

CASOS = Path(__file__).parent / "casos"


def bloques(*textos):
    return [Bloque(t, i) for i, t in enumerate(textos)]


def reglas(resultado):
    return {h.regla for h in resultado.hallazgos}


RELLENO = ("En el mundo actual, la automatización juega un papel crucial en las empresas pequeñas del país. "
           "Cabe señalar que su adopción crece cada año entre los negocios que venden por internet.")
MORALEJA = ("Las ventas subieron cuarenta por ciento en el trimestre después de responder en menos de un minuto. "
            "Los clientes nuevos llegaron sobre todo por recomendación de los antiguos clientes del local. "
            "Esto demuestra que responder rápido es la clave del crecimiento de cualquier negocio.")


# ------------------------------------------------------------------ lector
def test_frases_respeta_abreviaturas():
    t = "Según Wooldridge (2015, cap. 1) los datos importan. Sims et al. lo discuten en 1980. Fin del texto aquí."
    assert [f for f, _, _ in frases(t)] == [
        "Según Wooldridge (2015, cap. 1) los datos importan.",
        "Sims et al. lo discuten en 1980.",
        "Fin del texto aquí.",
    ]


def test_frases_corta_despues_de_un_año():
    t = "La caída fue en abril de 2020. Después llegó la reapertura por fases en todo el país."
    assert len(frases(t)) == 2


def test_markdown_separa_items_de_lista(tmp_path):
    md = tmp_path / "x.md"
    md.write_text("# Título\n\n- Primer elemento de la lista con bastante texto para contar.\n"
                  "- Segundo elemento de la lista con bastante texto para contar.\n\n## Referencias\n\n"
                  "Autor, A. (2020). Un libro que no debe revisarse; tiene punto y coma.\n", encoding="utf-8")
    b = leer(md)
    assert len(b) == 2
    assert all(";" not in x.texto for x in b)


def test_docx_salta_titulos_leyendas_y_referencias(tmp_path):
    d = docx.Document()
    d.add_heading("Introducción", level=1)
    d.add_paragraph("Este es un párrafo de prosa que sí se debe revisar porque tiene largo suficiente.")
    d.add_paragraph("Figura 1")
    d.add_paragraph("Nota. Elaboración propia con datos del BCRP y otras fuentes oficiales del país.")
    d.add_paragraph("Referencias")
    d.add_paragraph("Sims, C. A. (1980). Macroeconomics and reality; Econometrica, 48(1), 1-48, que no se revisa.")
    ruta = tmp_path / "x.docx"
    d.save(ruta)
    b = leer(ruta)
    assert len(b) == 1 and b[0].texto.startswith("Este es")


# ------------------------------------------------------------------ reglas
def test_relleno_es_regla_dura_en_todos_los_perfiles():
    for perfil in perfiles_disponibles():
        r = revisar(bloques(RELLENO), perfil)
        assert "relleno" in {h.regla for h in r.duras}, perfil


def test_punto_y_coma_solo_es_dura_en_estricto():
    t = ("El modelo es estable; sus raíces están dentro del círculo unitario y los residuos son limpios. "
         "La prueba LM no rechaza en ningún orden y la normalidad se mantiene en el límite del cinco por ciento.")
    assert "punto_y_coma" in {h.regla for h in revisar(bloques(t), "estricto").duras}
    assert "punto_y_coma" not in {h.regla for h in revisar(bloques(t), "academico").duras}


def test_punto_y_coma_entre_citas_no_cuenta():
    t = ("El error de medición sesga la elasticidad hacia cero (Nelson, 1994; Ravallion y Chaudhuri, 1997) y "
         "eso se ve en los datos de la encuesta nacional de hogares.")
    assert revisar(bloques(t), "estricto").metricas["punto_y_coma_x1000"] == 0


def test_rayas_en_estricto():
    t = "Muchas pymes dudan — por miedo o por desconocimiento — de dar el paso hacia la automatización completa."
    assert "rayas" in {h.regla for h in revisar(bloques(t), "estricto").duras}


def test_moraleja_y_cierre_en_narrativo():
    r = revisar(bloques(MORALEJA), "narrativo")
    assert {"moraleja", "cierre_moraleja"} <= reglas(r)


def test_dos_puntos_que_presentan_ecuacion_no_cuentan():
    t = "La ecuación estimada, que resume el modelo completo de la sección anterior, es la siguiente:"
    r = revisar(bloques(t), "estricto")
    assert r.metricas["dos_puntos_x1000"] == 0


def test_dos_puntos_de_etiqueta_no_cuentan():
    t = "Integrantes: Ana Pérez, Luis Gómez y María Torres, estudiantes del décimo ciclo de la carrera."
    assert revisar(bloques(t), "estricto").metricas["dos_puntos_x1000"] == 0


def test_molde_afirmacion_prueba_cuenta():
    t = ("La producción es persistente: su primer rezago tiene un coeficiente alto y significativo en el modelo. "
         "La inflación también: su inercia domina la ecuación y explica casi toda la variación mensual.")
    assert "dos_puntos" in {h.regla for h in revisar(bloques(t), "estricto").duras}


def test_tricolon_no_confunde_autores_ni_años():
    t = ("Para el Perú, Céspedes Reynaga y Huarancca (2025) estudian las olas 2023, 2024 y 2025 de la encuesta "
         "nacional de hogares con mucho detalle y cuidado metodológico.")
    assert revisar(bloques(t), "academico").metricas["tricolon_x1000"] == 0


def test_comillas_de_enfasis():
    t = "La \"transformación\" digital no espera a nadie y los negocios lo saben desde hace bastante tiempo."
    assert "comillas_enfasis" in reglas(revisar(bloques(t), "academico"))


def test_ritmo_uniforme():
    frase = "El modelo usa datos mensuales del banco central del país. "
    r = revisar(bloques(frase * 10), "academico")
    assert "ritmo_uniforme" in reglas(r)


def test_moraleja_apagada_en_academico():
    """La calibración mostró que en prosa académica humana estas fórmulas son más frecuentes que en la de IA."""
    assert "moraleja" not in reglas(revisar(bloques(MORALEJA), "academico"))


def test_senales_calibradas():
    sin_nada = ("Las ventas crecieron mucho, la empresa abrió locales nuevos y contrató gente, mientras la "
                "competencia, la regulación y el clima complicaban todo; además, el panorama se volvió integral, "
                "robusto, crucial y fundamental para el crecimiento sostenido de la región entera.")
    r = revisar(bloques(sin_nada * 3), "academico")
    assert {"Sin citas (Autor, año)", "Léxico de IA"} <= set(r.senales)


def test_senales_no_se_mezclan_entre_perfiles():
    redes, academico = cargar_perfil("redes"), cargar_perfil("academico")
    assert "citas_x1000" in academico["senales"] and "citas_x1000" not in redes["senales"]
    assert cargar_perfil("estricto")["senales"] == academico["senales"]
    assert cargar_perfil("narrativo")["senales"] == redes["senales"]


def test_perfil_hereda_y_sobrescribe():
    p = cargar_perfil("estricto")
    assert p["severidad"]["punto_y_coma"] == "dura"
    assert p["umbrales"]["moraleja"] == cargar_perfil("academico")["umbrales"]["moraleja"]


def test_perfil_propio_desde_archivo(tmp_path):
    propio = tmp_path / "mio.toml"
    propio.write_text('hereda = "academico"\n[severidad]\nrelleno = "apagada"\n', encoding="utf-8")
    r = revisar(bloques(RELLENO), str(propio))
    assert "relleno" not in reglas(r)


# ------------------------------------------------------------------ caso real
def test_marti_pasa_el_perfil_academico_y_falla_el_estricto():
    """Un texto humano de 1891: el punto y coma es norma del español, no huella de IA."""
    b = leer(CASOS / "marti_nuestra_america.md")
    assert revisar(b, "academico").pasa
    estricto = revisar(b, "estricto")
    # el estricto marca también sus comillas irónicas ("increíbles"): es gusto personal, no huella de IA
    assert {h.regla for h in estricto.duras} == {"punto_y_coma", "comillas_enfasis"}
    assert estricto.metricas["cv_largo_frase"] > 0.6  # ritmo humano: frases muy dispares


# ------------------------------------------------------------------ marcado
def test_marcar_docx_no_cambia_el_texto(tmp_path):
    d = docx.Document()
    p = d.add_paragraph("Primera parte del párrafo en ")
    p.add_run("negrita").bold = True
    p.add_run(". " + RELLENO)
    d.add_paragraph(MORALEJA)
    origen, destino = tmp_path / "a.docx", tmp_path / "b.docx"
    d.save(origen)
    r = revisar(leer(origen), "academico")
    n = marcar_docx(origen, destino, r)
    a, b = docx.Document(origen), docx.Document(destino)
    assert [x.text for x in a.paragraphs] == [x.text for x in b.paragraphs]
    assert n > 0 and len(list(b.comments)) == n
    assert any(run.bold for run in b.paragraphs[0].runs)  # el formato se conserva
    assert any(run.font.highlight_color is not None for x in b.paragraphs for run in x.runs)


def test_reporte_html_escapa_y_marca():
    b = bloques("Texto con <etiqueta> y " + RELLENO)
    html = reporte_html(b, revisar(b, "academico"))
    assert "&lt;etiqueta&gt;" in html and '<mark class="dura"' in html


# ------------------------------------------------------------------ cli
def test_cli_codigos_de_salida(tmp_path, capsys):
    malo, bueno = tmp_path / "malo.txt", tmp_path / "bueno.txt"
    malo.write_text(RELLENO, encoding="utf-8")
    bueno.write_text("El PBI cayó 29% en abril de 2020. Se recuperó en junio. La inflación siguió estable "
                     "durante todo ese periodo, sin saltos que llamaran la atención de nadie.", encoding="utf-8")
    assert main([str(malo)]) == 1
    assert main([str(bueno)]) == 0


def test_cli_perfil_desconocido(tmp_path):
    f = tmp_path / "x.txt"
    f.write_text(RELLENO, encoding="utf-8")
    with pytest.raises(ValueError):
        main([str(f), "--perfil", "inexistente"])
